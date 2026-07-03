import sys, os, subprocess
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = r"D:\project\java_project\skill_demo\target\skill-runtime-output\Spring_AI_Skill_Routing.docx"

# Tokens: standard_business_brief
def sr(r, s=11, c=None, b=None):
    r.font.name = "Calibri"; r.font.size = Pt(s)
    if c: r.font.color.rgb = RGBColor.from_string(c)
    if b is not None: r.bold = b
def sp(p, be=0, af=120, ln=264):
    pf = p.paragraph_format
    pf.space_before = Pt(be/20) if be else Pt(0)
    pf.space_after = Pt(af/20) if af else Pt(0)
    pf.line_spacing = ln/240
def hd(doc, txt, lv=1):
    p = doc.add_paragraph(); r = p.add_run(txt)
    sizes = {1:(16,"2E74B5",320,160),2:(13,"2E74B5",240,120),3:(12,"1F4D78",160,80)}
    s,c,b,a = sizes[lv]; sr(r, s, c, True); sp(p, b, a); p.style = doc.styles['Normal']
def bd(doc, txt):
    p = doc.add_paragraph(txt); p.style = doc.styles['Normal']; sp(p)
    for r in p.runs: sr(r)
def bu(doc, txt):
    p = doc.add_paragraph(txt); p.style = doc.styles['Normal']; sp(p, af=80)
    pf = p.paragraph_format; pf.left_indent = Inches(0.25); pf.first_line_indent = Inches(-0.25)
    for r in p.runs: sr(r)
def bl(doc, bpt, npt):
    p = doc.add_paragraph(); p.style = doc.styles['Normal']; sp(p)
    sr(p.add_run(bpt), b=True); sr(p.add_run(npt))
def sc(c, txt, b=False, cl=None):
    c.text = ""; p = c.paragraphs[0]; sr(p.add_run(txt), b=b, c=cl); sp(p, 0, 0, 240)
def shd(c, h):
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), h); sh.set(qn('w:val'), 'clear')
    c._tc.get_or_add_tcPr().append(sh)
def scm(c):
    tc_pr = c._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for s,v in (("top",80),("bottom",80),("start",120),("end",120)):
        e=OxmlElement(f'w:{s}'); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); m.append(e)
    tc_pr.append(m)
def ag(t, cw):
    t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.LEFT; tp = t._tbl.tblPr
    for tag,val in [("w:tblW",sum(cw)),("w:tblInd",120)]:
        e=OxmlElement(tag); e.set(qn('w:type'),'dxa'); e.set(qn('w:w'),str(int(val))); tp.append(e)
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tp.append(lay)
    g=t._tbl.tblGrid
    for ch in list(g): g.remove(ch)
    for w in cw:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(int(w))); g.append(gc)
    for row in t.rows:
        for i,c in enumerate(row.cells):
            c.width=Twips(cw[i]); tc_pr=c._tc.get_or_add_tcPr()
            tcw=OxmlElement('w:tcW'); tcw.set(qn('w:type'),'dxa'); tcw.set(qn('w:w'),str(int(cw[i]))); tc_pr.append(tcw); scm(c)

doc = Document()
s = doc.sections[0]
s.page_width = Inches(8.5); s.page_height = Inches(11)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
if doc.paragraphs: doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

st = doc.styles['Normal']
st.font.name = "Calibri"; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0,0,0)
st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.10

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Spring AI Skill Routing"); sr(r, 26); r.font.color.rgb = RGBColor(0,0,0); sp(p, 0, 60, 240)
p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(); r = p.add_run("A Concise Overview of Intelligent Request Dispatch"); sr(r, 12, "555555"); sp(p, 0, 160, 240)

hd(doc, "Introduction")
bd(doc, "Spring AI Skill Routing is a pattern for directing natural-language or structured requests to the most appropriate AI-driven handler \u2014 a \"skill\" \u2014 within a Spring-based application. Rather than hard-coding a single AI invocation path, a skill router introspects the incoming request, evaluates available skills, and dispatches execution to the best-matched handler.")
bd(doc, "This approach enables modular, extensible AI architectures where new capabilities can be added without rewriting existing orchestration logic. It is especially valuable in enterprise applications where different tasks require different models, prompts, tools, or vector stores.")

hd(doc, "Core Concepts")
hd(doc, "Skill", 2)
bd(doc, "A skill is a self-contained unit of AI functionality. Each skill defines a name, description, input schema, and an execution handler that may invoke a language model, run a retrieval chain, execute a tool, or call an external API. Skills are registered with the routing infrastructure at application startup.")
hd(doc, "Router", 2)
bd(doc, "The router is the central dispatch component. It receives a request \u2014 typically a user query or a structured intent \u2014 and matches it against the set of registered skills. Matching may use semantic similarity, keyword patterns, LLM-based classification, or a combination of strategies.")
hd(doc, "Routing Strategies", 2)
bu(doc, "Semantic routing \u2014 embed the request and find the skill with the closest description vector.")
bu(doc, "LLM-based routing \u2014 ask a language model to select the appropriate skill from a candidate list.")
bu(doc, "Keyword / regex routing \u2014 match predefined patterns to fixed skills.")
bu(doc, "Hybrid routing \u2014 combine multiple strategies with fallback logic for robustness.")

hd(doc, "Routing Strategy Comparison")
t = doc.add_table(rows=5, cols=4)
cw = [1800, 2700, 2700, 2160]
hdrs = ["Strategy", "Strength", "Limitation", "Best For"]
data = [
    ["Semantic", "Handles varied phrasing; no manual rules", "Requires embedding model; higher latency", "Open-ended Q&A, discovery"],
    ["LLM-based", "Flexible reasoning; context-aware", "Token cost; slower; model-dependent", "Complex or ambiguous intents"],
    ["Keyword / Regex", "Fast; deterministic; low cost", "Brittle; pattern maintenance overhead", "Well-defined commands, buttons"],
    ["Hybrid", "Balances speed and accuracy; fallback", "More moving parts; tuning effort", "Production systems with mixed traffic"],
]
for j, h in enumerate(hdrs):
    c = t.rows[0].cells[j]; sc(c, h, True, "FFFFFF"); shd(c, "2E74B5"); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
for i, rd in enumerate(data, 1):
    for j, v in enumerate(rd):
        c = t.rows[i].cells[j]; sc(c, v, j==0); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i % 2 == 0: shd(c, "F2F4F7")
ag(t, cw)

hd(doc, "Implementation Overview")
bd(doc, "In a typical Spring AI application, skill routing is wired through a central dispatcher bean. Skills are registered as Spring beans implementing a common interface. The router, itself a bean, collects all skill beans at injection time and exposes a dispatch method.")
bl(doc, "Key steps: ", "Define a Skill interface with canHandle(request) and execute(request) methods. Register concrete skill beans. Configure the router with a matching strategy. Expose a REST or message endpoint that delegates to the router.")
bd(doc, "Spring AI\u2019s auto-configuration and model abstraction layer make it straightforward to integrate skill routing with existing Spring Boot applications, supporting OpenAI, Azure OpenAI, Ollama, and other model providers interchangeably.")

hd(doc, "Key Benefits")
bu(doc, "Modularity \u2014 each skill is independently developed, tested, and deployed.")
bu(doc, "Extensibility \u2014 new skills plug in without altering existing dispatch logic.")
bu(doc, "Observability \u2014 routing decisions can be logged, traced, and monitored per skill.")
bu(doc, "Cost control \u2014 expensive LLM paths are reserved for requests that actually need them.")
bu(doc, "Graceful degradation \u2014 fallback skills handle unmatched requests instead of failing.")

hd(doc, "Conclusion")
bd(doc, "Spring AI Skill Routing provides a clean, extensible foundation for building AI-powered applications that can handle diverse request types through specialized handlers. By separating routing logic from skill execution, teams can evolve their AI capabilities incrementally while maintaining a consistent entry point for clients.")

doc.save(OUTPUT)
print(f"Created {OUTPUT}")

# Try to render
render_script = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "render_docx.py")
if os.path.exists(render_script):
    out_dir = os.path.join(os.path.dirname(OUTPUT), "render_out")
    os.makedirs(out_dir, exist_ok=True)
    res = subprocess.run([sys.executable, render_script, OUTPUT, "--output_dir", out_dir], capture_output=True, text=True, timeout=120)
    if res.returncode == 0 or os.listdir(out_dir):
        print("Render OK.")
        for f in sorted(os.listdir(out_dir)): print(f"  {f}")
    else:
        print("Render issues:", res.stderr[:500])
else:
    print("No render script")
